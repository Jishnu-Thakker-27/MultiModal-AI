import os
import logging
from pptx import Presentation
from typing import List, Dict, Any

logger = logging.getLogger("study_companion.ingestion.pptx")

def extract_pptx_content(file_path: str) -> List[Dict[str, Any]]:
    """
    Extracts text slide by slide from a PPTX file using python-pptx,
    preserving exact slide numbers.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"PPTX file not found at {file_path}")

    slides_content = []
    prs = Presentation(file_path)

    for idx, slide in enumerate(prs.slides):
        slide_text_elements = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    slide_text_elements.append(text)
        
        full_slide_text = "\n".join(slide_text_elements).strip()
        if full_slide_text:
            slides_content.append({
                "slide_number": idx + 1,
                "text": full_slide_text
            })

    logger.info(f"Extracted {len(slides_content)} slides from PPTX: {os.path.basename(file_path)}")
    return slides_content
